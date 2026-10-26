"""Sinh implement/User_Stories_IP_Camera_Web.md và implement/Product_Backlog_IP_Camera_Web.xlsx từ backlog_data.py."""
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).parent))
from backlog_data import DOD, DOR, EPICS, ITEMS, SPRINTS  # noqa: E402

OUT = Path(sys.argv[1])
OUT.mkdir(parents=True, exist_ok=True)
CAPACITY = 30
EPIC_NAME = dict(EPICS)
SPRINT_ORDER = [s[0] for s in SPRINTS]
MILESTONE = {s[0]: s[1] for s in SPRINTS}
BY_ID = {i["id"]: i for i in ITEMS}

# ---------- Kiểm tra dữ liệu ----------
errors = []
for it in ITEMS:
    for d in [x.strip() for x in it["deps"].split(",")]:
        if d in BY_ID and SPRINT_ORDER.index(BY_ID[d]["sprint"]) > SPRINT_ORDER.index(it["sprint"]):
            errors.append(f"{it['id']} ({it['sprint']}) phụ thuộc {d} ({BY_ID[d]['sprint']}) ở sprint sau")
    if it["pts"] not in (1, 2, 3, 5, 8):
        errors.append(f"{it['id']} điểm {it['pts']} không thuộc thang Fibonacci ≤ 8")
us_ids = sorted(i["id"] for i in ITEMS if i["type"] == "Story")
assert us_ids == [f"US-{n:02d}" for n in range(1, 45)], "Thiếu hoặc thừa US"
if errors:
    print("\n".join(errors)); sys.exit(1)

sprint_pts = defaultdict(int)
for it in ITEMS:
    sprint_pts[it["sprint"]] += it["pts"]
total = sum(i["pts"] for i in ITEMS)


def gherkin(it):
    out = []
    for name, lines in it["scenarios"]:
        out.append(f"Kịch bản: {name}")
        out += [f"  {ln}" for ln in lines]
    return "\n".join(out)


# ---------- Markdown ----------
md = []
md.append("# User Story & Product Backlog – CamHub: Nền tảng Web quản lý Camera IP\n")
md.append("> Ngày: 23/10/2026 · Người soạn: Product Owner / Business Analyst · Phiên bản: 0.1 (Bản nháp)")
md.append("> Đầu vào: [PRD](../requirements/PRD_IP_Camera_Web.md) v0.4 · [Use case](../requirements/Use_Case_IP_Camera_Web.md) v0.2 · "
          "[SADD](../requirements/SADD_IP_Camera_Web.md) v0.1 · [HTML design](../design/index.html) v0.1")
md.append("> Bảng theo dõi: [Product_Backlog_IP_Camera_Web.xlsx](Product_Backlog_IP_Camera_Web.xlsx) (cùng dữ liệu, có cột trạng thái và công thức tổng hợp sprint)\n")

md.append("## 1. Tóm tắt\n")
md.append(f"Backlog phiên bản 1.0 gồm **{len(us_ids)} user story** (từ PRD) và **{len(ITEMS) - len(us_ids)} enabler kỹ thuật** (từ SADD), "
          f"tổng **{total} điểm**. Với sprint 2 tuần và năng lực khoảng {CAPACITY} điểm, backlog cần **{len(SPRINTS)} sprint (20 tuần)**, "
          f"gồm 1 sprint PoC và nền tảng (S0) và 9 sprint tính năng.\n")
md.append("**Điểm khác với kế hoạch trong PRD (mục 10):** tổng thời gian vẫn là 20 tuần (PoC 2 tuần + 18 tuần), nhưng ranh giới các mốc dịch chuyển:\n")
md.append("| Mốc | PRD | Theo ước lượng backlog | Lý do |\n| --- | --- | --- | --- |")
md.append("| M1 – Alpha (toàn bộ Must) | 8 tuần (+2 tuần PoC) | 12 tuần: S0–S5 | Story Must cùng enabler cần khoảng 175 điểm, vượt 120 điểm của 4 sprint |")
md.append("| M2 – Beta pilot (Should) | 6 tuần | 6 tuần: S6–S8 | Không đổi |")
md.append("| M3 – GA (Could, kiểm thử, UAT) | 4 tuần | 2 tuần: S9 | Chỉ còn khoảng 14 điểm dự phòng cho sửa lỗi pilot và UAT |\n")
md.append("Rủi ro: M3 còn ít dự phòng. Nên cân nhắc thêm 1 sprint (S10) cho UAT và sửa lỗi, hoặc chuyển US-15 (Could) sang phiên bản sau (xem mục 8).\n")

md.append("## 2. Quy ước\n")
md.append("- **Mẫu story:** \"Là [vai trò], tôi muốn [hành động], để [giá trị]\"; đáp ứng INVEST (độc lập, thương lượng được, có giá trị, ước lượng được, nhỏ, kiểm thử được).")
md.append("- **Tiêu chí chấp nhận:** viết theo Gherkin tiếng Việt (*Kịch bản / Giả sử / Khi / Thì / Và*), mỗi kịch bản là một test tự động.")
md.append("- **Story point:** thang Fibonacci 1, 2, 3, 5, 8. Story lớn hơn 8 phải tách.")
md.append("- **Ưu tiên:** MoSCoW lấy từ BRS/PRD. Phiên bản 1.0 làm toàn bộ; ưu tiên dùng để xếp thứ tự sprint.")
md.append("- **Enabler (TS-xx):** công việc kỹ thuật không trực tiếp tạo tính năng, lấy từ SADD (PoC, khung dự án, hạ tầng).")
md.append("- **Truy vết:** mỗi item ghi BR (BRS), UC (Use case), màn hình (design/) và thành phần (SADD).\n")
md.append("### Definition of Ready\n")
md += [f"- [ ] {x}" for x in DOR]
md.append("\n### Definition of Done\n")
md += [f"- [ ] {x}" for x in DOD]

md.append("\n## 3. Kế hoạch sprint\n")
md.append("| Sprint | Mốc | Mục tiêu sprint | Item | Điểm |\n| --- | --- | --- | --- | --- |")
for s, ms, goal in SPRINTS:
    ids = [i["id"] for i in ITEMS if i["sprint"] == s]
    md.append(f"| {s} | {ms} | {goal} | {', '.join(ids)} | {sprint_pts[s]} / {CAPACITY} |")
md.append(f"| **Tổng** | | | {len(ITEMS)} item | **{total} / {CAPACITY * len(SPRINTS)}** |\n")

md.append("## 4. Epic\n")
md.append("| Epic | Tên | Số item | Điểm | Must | Should | Could |\n| --- | --- | --- | --- | --- | --- | --- |")
for eid, name in EPICS:
    its = [i for i in ITEMS if i["epic"] == eid]
    p = lambda pr: sum(i["pts"] for i in its if i["prio"] == pr)  # noqa: E731
    md.append(f"| {eid} | {name} | {len(its)} | {sum(i['pts'] for i in its)} | {p('Must')} | {p('Should')} | {p('Could')} |")

md.append("\n## 5. Chi tiết backlog\n")
for eid, name in EPICS:
    md.append(f"### {eid} {name}\n")
    for it in [i for i in ITEMS if i["epic"] == eid]:
        md.append(f"#### {it['id']} {it['title']}\n")
        md.append(f"**Là {it['persona']}, tôi muốn {it['want']}, để {it['so']}.**\n")
        md.append("| Ưu tiên | Điểm | Sprint | Phụ thuộc | BR | Use case | Màn hình | Thành phần (SADD) |\n| --- | --- | --- | --- | --- | --- | --- | --- |")
        md.append(f"| {it['prio']} | {it['pts']} | {it['sprint']} | {it['deps']} | {it['br']} | {it['uc']} | {it['screen']} | {it['sadd']} |\n")
        md.append("```gherkin")
        md.append(f"Tính năng: {it['title']}")
        for name_, lines in it["scenarios"]:
            md.append(f"\n  Kịch bản: {name_}")
            md += [f"    {ln}" for ln in lines]
        md.append("```\n")
        md.append(f"*Ghi chú kỹ thuật:* {it['notes']}\n")

md.append("## 6. Ma trận truy vết\n")
md.append("| Item | BR | Use case | Màn hình | Sprint |\n| --- | --- | --- | --- | --- |")
for it in ITEMS:
    md.append(f"| {it['id']} | {it['br']} | {it['uc']} | {it['screen']} | {it['sprint']} |")

md.append("\n## 7. Giả định\n")
md.append("- Đội khoảng 5–6 người (2 backend Java, 1 Python, 2 frontend Vue, 1 QA) với vận tốc khoảng 30 điểm/sprint; cần hiệu chỉnh sau 2 sprint đầu.")
md.append("- Story point do BA ước lượng ban đầu; đội phát triển cần ước lượng lại trong buổi refinement (ví dụ planning poker).")
md.append("- Sản phẩm tự phát triển (OQ-02 chưa chốt chính thức).\n")
md.append("## 8. Câu hỏi mở\n")
md.append("| Mã | Câu hỏi | Ảnh hưởng |\n| --- | --- | --- |")
md.append("| BQ-01 | Chấp nhận M1 kéo dài thành 12 tuần (S0–S5) không? | Lịch demo Alpha |")
md.append("| BQ-02 | Thêm S10 cho UAT/sửa lỗi, hay chuyển US-15 (Could) sang phiên bản sau? | Ngày GA |")
md.append("| BQ-03 | Zalo OA đã sẵn sàng chưa (OQ-P02)? | US-36 ở S7 |")
md.append("| BQ-04 | Thời gian lưu mặc định (OQ-04) | Giá trị mặc định của US-21 |")
(OUT / "User_Stories_IP_Camera_Web.md").write_text("\n".join(md) + "\n", encoding="utf-8")

# ---------- Excel ----------
F = "Arial"
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
hdr_fill = PatternFill("solid", start_color="1F3864")
hdr_font = Font(name=F, bold=True, color="FFFFFF", size=10)
body = Font(name=F, size=10)
bold = Font(name=F, size=10, bold=True)
input_fill = PatternFill("solid", start_color="FFF2CC")
total_fill = PatternFill("solid", start_color="D9E1F2")
wrap = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="top", wrap_text=True)
STATUSES = ["To Do", "In Progress", "In Review", "Done", "Blocked"]


def header(ws, row, cols, widths):
    for c, (h, w) in enumerate(zip(cols, widths), start=1):
        cell = ws.cell(row, c, h)
        cell.font, cell.fill, cell.alignment, cell.border = hdr_font, hdr_fill, center, border
        ws.column_dimensions[cell.column_letter].width = w


wb = Workbook()
g = wb.active
g.title = "Hướng dẫn"
g.sheet_view.showGridLines = False
g.column_dimensions["A"].width = 24
g.column_dimensions["B"].width = 100
guide = [
    ("Product Backlog – CamHub v1.0", None),
    ("Phiên bản", "0.1 · 23/10/2026 · nguồn chi tiết: User_Stories_IP_Camera_Web.md"),
    (None, None),
    ("Cách dùng", None),
    ("1", "Sheet \"Backlog\": mỗi dòng là một enabler (TS) hoặc user story (US). Cập nhật các ô nền vàng: Điểm, Sprint, Trạng thái, Người phụ trách."),
    ("2", "Cột \"Mốc\" tự tra theo Sprint từ sheet \"Kế hoạch sprint\"; không nhập tay."),
    ("3", "Sheet \"Kế hoạch sprint\": đổi năng lực mặc định ở ô B2 (hiện 30 điểm). Cột tỷ lệ tô đỏ khi sprint vượt năng lực."),
    ("4", "Sheet \"Epic\" và \"Mốc\" tự tổng hợp từ Backlog bằng công thức."),
    (None, None),
    ("Trạng thái", None),
    ("To Do", "Chưa bắt đầu"), ("In Progress", "Đang làm"), ("In Review", "Đang review code / chờ PO chấp nhận"),
    ("Done", "Đạt Definition of Done"), ("Blocked", "Bị chặn; ghi lý do ở cột Ghi chú"),
    (None, None),
    ("Definition of Ready", None),
] + [("•", x) for x in DOR] + [(None, None), ("Definition of Done", None)] + [("•", x) for x in DOD] + [
    (None, None),
    ("Giả định", "Năng lực 30 điểm/sprint 2 tuần; điểm do BA ước lượng ban đầu, đội phát triển cần ước lượng lại."),
]
for r, (a, b) in enumerate(guide, start=1):
    ca, cb = g.cell(r, 1, a), g.cell(r, 2, b)
    ca.font = cb.font = body
    ca.alignment = cb.alignment = wrap
    if b is None and a:
        ca.font = Font(name=F, size=14 if r == 1 else 11, bold=True, color="1F3864")

# Backlog
bl = wb.create_sheet("Backlog")
cols = ["ID", "Loại", "Epic", "Tiêu đề", "User story", "Ưu tiên", "Điểm", "Sprint", "Mốc", "Trạng thái",
        "Phụ thuộc", "BR", "Use case", "Màn hình", "Thành phần SADD", "Tiêu chí chấp nhận (Gherkin)", "Người phụ trách", "Ghi chú kỹ thuật"]
widths = [8, 9, 26, 32, 50, 9, 7, 8, 18, 12, 18, 16, 16, 26, 30, 70, 16, 45]
header(bl, 1, cols, widths)
bl.row_dimensions[1].height = 30
for r, it in enumerate(ITEMS, start=2):
    vals = [it["id"], it["type"], f"{it['epic']} {EPIC_NAME[it['epic']]}", it["title"],
            f"Là {it['persona']}, tôi muốn {it['want']}, để {it['so']}.", it["prio"], it["pts"], it["sprint"],
            f"=IFERROR(VLOOKUP(H{r},'Kế hoạch sprint'!$A$5:$B${4 + len(SPRINTS)},2,FALSE),\"\")",
            "To Do", it["deps"], it["br"], it["uc"], it["screen"], it["sadd"], gherkin(it), None, it["notes"]]
    for c, v in enumerate(vals, start=1):
        cell = bl.cell(r, c, v)
        cell.font, cell.border = body, border
        cell.alignment = center if c in (1, 2, 6, 7, 8, 10) else wrap
        if c in (7, 8, 10, 17):
            cell.fill = input_fill
    bl.cell(r, 1).font = bold
last = len(ITEMS) + 1
bl.freeze_panes = "E2"
bl.auto_filter.ref = f"A1:R{last}"
dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=False)
bl.add_data_validation(dv); dv.add(f"J2:J{last}")
dvs = DataValidation(type="list", formula1='"' + ",".join(SPRINT_ORDER + ["Backlog"]) + '"', allow_blank=True)
bl.add_data_validation(dvs); dvs.add(f"H2:H{last}")
dvp = DataValidation(type="list", formula1='"1,2,3,5,8,13"', allow_blank=True)
bl.add_data_validation(dvp); dvp.add(f"G2:G{last}")
for s, (bg, fg) in {"Done": ("C6EFCE", "006100"), "In Progress": ("DDEBF7", "1F4E78"), "In Review": ("FFEB9C", "9C5700"), "Blocked": ("FFC7CE", "9C0006")}.items():
    bl.conditional_formatting.add(f"J2:J{last}", CellIsRule(operator="equal", formula=[f'"{s}"'],
                                  fill=PatternFill("solid", start_color=bg), font=Font(name=F, color=fg, bold=True)))

BLG = f"Backlog!$G$2:$G${last}"; BLH = f"Backlog!$H$2:$H${last}"; BLJ = f"Backlog!$J$2:$J${last}"
BLC = f"Backlog!$C$2:$C${last}"; BLF = f"Backlog!$F$2:$F${last}"; BLI = f"Backlog!$I$2:$I${last}"

# Kế hoạch sprint
sp = wb.create_sheet("Kế hoạch sprint")
sp.sheet_view.showGridLines = False
sp["A1"] = "Kế hoạch sprint (2 tuần/sprint)"; sp["A1"].font = Font(name=F, size=14, bold=True, color="1F3864")
sp["A2"] = "Năng lực mặc định (điểm/sprint)"; sp["A2"].font = bold
sp["B2"] = CAPACITY; sp["B2"].fill = input_fill; sp["B2"].font = Font(name=F, size=10, color="0000FF")
sp["C2"] = "Nguồn: giả định đội 5–6 người (đã chốt với người dùng 23/10/2026)"; sp["C2"].font = Font(name=F, size=9, italic=True, color="595959")
spc = ["Sprint", "Mốc", "Mục tiêu sprint", "Năng lực", "Điểm kế hoạch", "Tỷ lệ sử dụng", "Số item", "Điểm Done", "Điểm còn lại"]
header(sp, 4, spc, [8, 18, 60, 10, 13, 13, 9, 11, 12])
for i, (s, ms, goal) in enumerate(SPRINTS):
    r = 5 + i
    row = [s, ms, goal, "=$B$2", f"=SUMIFS({BLG},{BLH},A{r})", f'=IF(D{r}=0,"-",E{r}/D{r})',
           f"=COUNTIFS({BLH},A{r})", f'=SUMIFS({BLG},{BLH},A{r},{BLJ},"Done")', f"=E{r}-H{r}"]
    for c, v in enumerate(row, start=1):
        cell = sp.cell(r, c, v); cell.font, cell.border = body, border
        cell.alignment = wrap if c == 3 else Alignment(horizontal="center", vertical="top")
    sp.cell(r, 4).fill = input_fill
    sp.cell(r, 6).number_format = "0%"
tr = 5 + len(SPRINTS)
sp.cell(tr, 1, "Tổng")
for c, L in [(4, "D"), (5, "E"), (7, "G"), (8, "H"), (9, "I")]:
    sp.cell(tr, c, f"=SUM({L}5:{L}{tr - 1})")
sp.cell(tr, 6, f'=IF(D{tr}=0,"-",E{tr}/D{tr})'); sp.cell(tr, 6).number_format = "0%"
for c in range(1, 10):
    cell = sp.cell(tr, c); cell.font, cell.fill, cell.border = bold, total_fill, border
    cell.alignment = Alignment(horizontal="center")
sp.cell(tr + 1, 1, "Kiểm tra: điểm của item chưa xếp sprint").font = Font(name=F, size=9, italic=True, color="595959")
sp.cell(tr + 1, 5, f"=SUM({BLG})-E{tr}").font = body
sp.conditional_formatting.add(f"F5:F{tr - 1}", CellIsRule(operator="greaterThan", formula=["1"],
                              fill=PatternFill("solid", start_color="FFC7CE"), font=Font(name=F, color="9C0006", bold=True)))

# Epic
ep = wb.create_sheet("Epic")
ep.sheet_view.showGridLines = False
header(ep, 1, ["Epic", "Số item", "Tổng điểm", "Must", "Should", "Could", "Điểm Done", "Tiến độ"], [40, 9, 11, 9, 9, 9, 11, 10])
for i, (eid, name) in enumerate(EPICS):
    r = 2 + i
    key = f"{eid} {name}"
    row = [key, f"=COUNTIFS({BLC},A{r})", f"=SUMIFS({BLG},{BLC},A{r})", f'=SUMIFS({BLG},{BLC},A{r},{BLF},"Must")',
           f'=SUMIFS({BLG},{BLC},A{r},{BLF},"Should")', f'=SUMIFS({BLG},{BLC},A{r},{BLF},"Could")',
           f'=SUMIFS({BLG},{BLC},A{r},{BLJ},"Done")', f'=IF(C{r}=0,"-",G{r}/C{r})']
    for c, v in enumerate(row, start=1):
        cell = ep.cell(r, c, v); cell.font, cell.border = body, border
        cell.alignment = Alignment(horizontal="left" if c == 1 else "center")
    ep.cell(r, 8).number_format = "0%"
er = 2 + len(EPICS)
ep.cell(er, 1, "Tổng")
for c, L in [(2, "B"), (3, "C"), (4, "D"), (5, "E"), (6, "F"), (7, "G")]:
    ep.cell(er, c, f"=SUM({L}2:{L}{er - 1})")
ep.cell(er, 8, f'=IF(C{er}=0,"-",G{er}/C{er})'); ep.cell(er, 8).number_format = "0%"
for c in range(1, 9):
    cell = ep.cell(er, c); cell.font, cell.fill, cell.border = bold, total_fill, border

# Mốc
mo = wb.create_sheet("Mốc")
mo.sheet_view.showGridLines = False
header(mo, 1, ["Mốc", "Số sprint", "Số tuần", "Điểm kế hoạch", "Năng lực", "Điểm Done", "Tiến độ"], [22, 10, 9, 14, 10, 11, 10])
mss = list(dict.fromkeys(MILESTONE.values()))
SPB = f"'Kế hoạch sprint'!$B$5:$B${tr - 1}"; SPD = f"'Kế hoạch sprint'!$D$5:$D${tr - 1}"
for i, m in enumerate(mss):
    r = 2 + i
    row = [m, f"=COUNTIFS({SPB},A{r})", f"=B{r}*2", f"=SUMIFS({BLG},{BLI},A{r})", f"=SUMIFS({SPD},{SPB},A{r})",
           f'=SUMIFS({BLG},{BLI},A{r},{BLJ},"Done")', f'=IF(D{r}=0,"-",F{r}/D{r})']
    for c, v in enumerate(row, start=1):
        cell = mo.cell(r, c, v); cell.font, cell.border = body, border
        cell.alignment = Alignment(horizontal="left" if c == 1 else "center")
    mo.cell(r, 7).number_format = "0%"

wb.move_sheet("Kế hoạch sprint", offset=-1)
wb.active = 1
wb.calculation.fullCalcOnLoad = True
wb.save(OUT / "Product_Backlog_IP_Camera_Web.xlsx")

print("total", total, "items", len(ITEMS))
for s in SPRINT_ORDER:
    print(s, MILESTONE[s], sprint_pts[s])
